#!/usr/bin/env python3
"""Generate all MAGNETO manuscript and supplementary tables.

Runs each table script in scripts/tables, merges the DOCX files into
main_tables.docx and supplementary_tables.docx, and writes a manifest and
validation report.
"""
from __future__ import annotations

import argparse
import csv
import subprocess
import sys
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parent.parent
TABLES_DIR = PROJECT_ROOT / "scripts" / "tables"
REPORTS_TABLES_DIR = PROJECT_ROOT / "reports" / "tables"

MAIN_TABLES = ["table1_primary_results"]
SUPP_TABLES = [f"tableS{i}" for i in range(1, 11)]
SUPP_TABLES = [
    "tableS1_data_sources_qc",
    "tableS2_sample_definitions",
    "tableS3_fixed_window_results",
    "tableS4_temporal_surrogate",
    "tableS5_kp_f107",
    "tableS6_environmental_driver",
    "tableS7_shapley",
    "tableS8_landcover_pooled",
    "tableS9_lai_diagnostics",
    "tableS10_sif_wavelength",
]
ALL_TABLES = MAIN_TABLES + SUPP_TABLES


def run_script(script: Path) -> tuple[bool, str]:
    """Run one table script and return (ok, message)."""
    try:
        result = subprocess.run(
            [sys.executable, str(script)],
            cwd=PROJECT_ROOT,
            env={**dict(subprocess.os.environ), "PYTHONPATH": str(PROJECT_ROOT / "scripts")},
            capture_output=True,
            text=True,
        )
    except Exception as exc:  # pragma: no cover
        return False, f"launcher error: {exc}"
    if result.returncode != 0:
        return False, result.stderr.strip() or result.stdout.strip()
    return True, result.stdout.strip().splitlines()[-1]


def merge_docx(parts: list[Path], output: Path) -> None:
    """Merge a list of DOCX files into one document using docxcompose."""
    from docx import Document
    from docxcompose.composer import Composer

    if not parts or not parts[0].exists():
        Document().save(output)
        return

    merged = Document(parts[0])
    composer = Composer(merged)
    for part in parts[1:]:
        if part.exists():
            composer.append(Document(part))
    composer.save(output)


def build_manifest(rows: list[dict]) -> None:
    """Write tables_manifest.csv."""
    REPORTS_TABLES_DIR.mkdir(parents=True, exist_ok=True)
    path = REPORTS_TABLES_DIR / "tables_manifest.csv"
    fieldnames = [
        "table_id", "title", "source_csv", "markdown", "docx",
        "row_count", "input_files", "status", "notes",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"[OK] Manifest: {path}")


def build_validation(rows: list[dict]) -> None:
    """Write tables_validation.md."""
    REPORTS_TABLES_DIR.mkdir(parents=True, exist_ok=True)
    path = REPORTS_TABLES_DIR / "tables_validation.md"
    lines = ["# Tables Validation Report", ""]
    for r in rows:
        lines.append(f"## {r['table_id']}")
        lines.append(f"- **Title:** {r['title']}")
        lines.append(f"- **Status:** {r['status']}")
        lines.append(f"- **Rows:** {r['row_count']}")
        lines.append(f"- **Source CSV:** {r['source_csv']}")
        lines.append(f"- **Inputs:** {r['input_files']}")
        if r.get("notes"):
            lines.append(f"- **Notes:** {r['notes']}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    print(f"[OK] Validation report: {path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate MAGNETO tables.")
    parser.add_argument("--tables", type=str, default="", help="Comma-separated table basenames to run.")
    args = parser.parse_args()

    tables = args.tables.split(",") if args.tables else ALL_TABLES
    tables = [t.strip() for t in tables if t.strip()]

    REPORTS_TABLES_DIR.mkdir(parents=True, exist_ok=True)

    rows = []
    failures = []
    main_docx_parts = []
    supp_docx_parts = []

    for name in tables:
        script = TABLES_DIR / f"{name}.py"
        print(f"Running {name} ...", end=" ", flush=True)
        ok, msg = run_script(script)
        print("OK" if ok else f"FAILED: {msg}")
        if not ok:
            failures.append((name, msg))
            continue

        source_csv = REPORTS_TABLES_DIR / f"{name}_source.csv"
        md = REPORTS_TABLES_DIR / f"{name}.md"
        docx = REPORTS_TABLES_DIR / f"{name}.docx"

        row_count = 0
        if source_csv.exists():
            with open(source_csv, encoding="utf-8") as f:
                row_count = sum(1 for _ in f) - 1

        rows.append({
            "table_id": name,
            "title": " ".join(name.replace("table", "Table ").replace("S", " S").split("_")),
            "source_csv": str(source_csv.relative_to(PROJECT_ROOT)),
            "markdown": str(md.relative_to(PROJECT_ROOT)),
            "docx": str(docx.relative_to(PROJECT_ROOT)),
            "row_count": row_count,
            "input_files": "see script docstring",
            "status": "complete" if docx.exists() else "failed",
            "notes": msg,
        })

        if name in MAIN_TABLES:
            main_docx_parts.append(docx)
        else:
            supp_docx_parts.append(docx)

    if failures:
        print("\nFailures:")
        for name, msg in failures:
            print(f"  - {name}: {msg}")
        return 1

    merge_docx(main_docx_parts, REPORTS_TABLES_DIR / "main_tables.docx")
    merge_docx(supp_docx_parts, REPORTS_TABLES_DIR / "supplementary_tables.docx")

    build_manifest(rows)
    build_validation(rows)

    print("\nAll tables generated.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
