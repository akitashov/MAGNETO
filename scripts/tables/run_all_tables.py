#!/usr/bin/env python3
"""Orchestrate generation of all manuscript and supplementary tables.

Rebuilds every table from existing analytical outputs, merges DOCX files, writes
a manifest and a validation report, and runs pytest.
"""
from __future__ import annotations

import hashlib
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

# Allow imports of sibling table scripts and shared utilities.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docxcompose.composer import Composer

from tables._table_utils import project_root, tables_dir, write_table
from tables.table1_primary_results import build_table1
from tables.tableS1_data_sources_qc import build_table_s1
from tables.tableS2_sample_definitions import build_table_s2
from tables.tableS3_fixed_window_results import build_table_s3
from tables.tableS4_temporal_surrogate import build_table_s4
from tables.tableS5_kp_f107 import build_table_s5
from tables.tableS6_environmental_driver import build_table_s6
from tables.tableS7_shapley import build_table_s7
from tables.tableS8_landcover_pooled import build_table_s8
from tables.tableS9_lai_diagnostics import build_table_s9
from tables.tableS10_sif_wavelength import build_table_s10


TABLE_SPECS: list[dict] = [
    {
        "name": "table1_primary_results",
        "caption": "Table 1. Primary SII–SIF associations and key robustness analyses",
        "landscape": True,
        "build": build_table1,
        "inputs": [
            "results/fixed_window_results.csv",
            "results/surrogate_summary.csv",
            "results/supplementary_checks/table1_bootstrap_summary.csv",
        ],
        "filters": [
            "sample_type in pooled_full, pairwise_matched, control_vegetated, control_Sahara",
            "method = harmonic, except pairwise matched = cyclic_spline",
            "sii_window = sii_28d, except one sensitivity row = sii_21d",
        ],
        "p_value_source": "results/surrogate_summary.csv (empirical temporal-surrogate, B = 1000)",
        "bootstrap_source": "results/supplementary_checks/table1_bootstrap_summary.csv (cell-cluster bootstrap, B = 1000)",
        "sample_matching": "Pairwise-matched row uses identical observations for harmonic vs cyclic-spline comparison",
    },
    {
        "name": "tableS1_data_sources_qc",
        "caption": "Table S1. Data sources, spatial and temporal resolution, preprocessing and QC",
        "landscape": True,
        "build": build_table_s1,
        "inputs": [
            "results/qc_flow.csv",
            "reports/SIF_QC_AUDIT.md",
        ],
        "filters": ["Assembled from project documentation and QC manifests"],
        "p_value_source": "N/A",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS2_sample_definitions",
        "caption": "Table S2. Analytical sample and scenario definitions",
        "landscape": False,
        "build": build_table_s2,
        "inputs": [
            "results/control_definitions.json",
            "results/fixed_window_results.csv",
            "results/supplementary_checks/landcover_analysis.csv",
            "data/processed/environmental_driver_input.parquet",
        ],
        "filters": [
            "harmonic method, sii_28d window",
            "SAA counts derived from is_SAA flag in environmental_driver_input.parquet",
            "land-cover classes from landcover_analysis.csv sample_type = landcover_*",
        ],
        "p_value_source": "N/A",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS3_fixed_window_results",
        "caption": "Table S3. Fixed-window SII–SIF results by scenario, temperature regime and detrending method",
        "landscape": True,
        "build": build_table_s3,
        "inputs": [
            "results/supplementary_checks/temperature_scenarios_harmonic.csv",
            "results/supplementary_checks/temperature_scenarios_spline.csv",
        ],
        "filters": [
            "All scenarios in temperature_scenarios CSVs plus explicit Sahara × Frozen not-estimable rows (1 cell, 3 observations) from data/processed/environmental_driver_input.parquet",
            "window_days = 21 and 28",
            "Inferential values suppressed where n_cells < 50",
        ],
        "p_value_source": "temperature_scenarios CSVs (empirical temporal-surrogate, B = 1000)",
        "bootstrap_source": "temperature_scenarios CSVs (ci_low, ci_high)",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS4_temporal_surrogate",
        "caption": "Table S4. Temporal-surrogate inference for the primary analysis rows",
        "landscape": True,
        "build": build_table_s4,
        "inputs": [
            "results/surrogate_summary.csv",
        ],
        "filters": [
            "sample_type in pooled_full, pairwise_matched, control_vegetated, control_strict_low_lai, control_Sahara (control_SAA surrogate rows are absent from surrogate_summary.csv)",
            "All three surrogate modes",
            "Sample basis distinguishes native method-specific samples from the pairwise-matched sample",
        ],
        "p_value_source": "results/surrogate_summary.csv (empirical plus-one p, B = 1000)",
        "bootstrap_source": "N/A",
        "sample_matching": "Pairwise-matched basis identified in separate Sample basis column",
    },
    {
        "name": "tableS5_kp_f107",
        "caption": "Table S5. Alternative geomagnetic and solar-activity comparisons",
        "landscape": True,
        "build": build_table_s5,
        "inputs": [
            "results/supplementary_checks/kp_comparison.csv",
            "results/supplementary_checks/f107_comparison.csv",
            "results/supplementary_checks/saa_control.csv",
        ],
        "filters": [
            "sample_type = pooled_sii, pooled_kp, pooled_f10_7, control_SAA",
            "method = harmonic",
            "sii_window_days = 28",
        ],
        "p_value_source": "Nominal Spearman p-values from comparison CSVs (descriptive only)",
        "bootstrap_source": "N/A",
        "sample_matching": "F10.7 comparison uses matched SII reference from f107_comparison.csv; Kp comparison uses full pooled SII reference; exact n_obs read from source CSVs",
    },
    {
        "name": "tableS6_environmental_driver",
        "caption": "Table S6. Environmental-driver adjustment summary across PAR × VPD specifications",
        "landscape": True,
        "build": build_table_s6,
        "inputs": [
            "results/environmental_driver_matrix_full_summary.csv",
        ],
        "filters": [
            "scenario in full_sample, persistently_vegetated",
            "All six temperature bins and SII windows 1–28 days",
            "DOCX view restricted to windows 1, 7, 14, 21, 28",
        ],
        "p_value_source": "N/A",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS7_shapley",
        "caption": "Table S7. Shapley/LMG decomposition of residual SIF variance at the 28-day common window",
        "landscape": True,
        "build": build_table_s7,
        "inputs": [
            "results/supplementary_checks/driver_r2_shapley_by_window.csv",
        ],
        "filters": [
            "window_days = 28",
            "All five scenarios and six temperature bins",
        ],
        "p_value_source": "N/A",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS8_landcover_pooled",
        "caption": "Table S8. Pooled land-cover associations with SII",
        "landscape": True,
        "build": build_table_s8,
        "inputs": [
            "results/supplementary_checks/landcover_analysis.csv",
            "results/supplementary_checks/landcover_analysis_surrogates.csv",
        ],
        "filters": [
            "method = harmonic, cyclic_spline",
            "sii_window = sii_28d",
            "sample_type = landcover_* (excluding landcover_temp_*)",
        ],
        "p_value_source": "landcover_analysis_surrogates.csv (empirical temporal-surrogate, B = 1000)",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS9_lai_diagnostics",
        "caption": "Table S9. LAI-stratification of SII–SIF associations",
        "landscape": True,
        "build": build_table_s9,
        "inputs": [
            "results/control_definitions.json",
            "results/fixed_window_results.csv",
            "results/surrogate_summary.csv",
            "data/interim/lai_quartile_boundaries.json",
            "results/supplementary_checks/lai_quartile_bootstrap_summary.csv",
        ],
        "filters": [
            "sample_type = control_vegetated, lai_quartile",
            "method = harmonic",
            "sii_window = sii_28d",
        ],
        "p_value_source": "results/surrogate_summary.csv (empirical temporal-surrogate, B = 1000)",
        "bootstrap_source": "results/supplementary_checks/lai_quartile_bootstrap_summary.csv (cell-cluster bootstrap, B = 1000)",
        "sample_matching": "N/A",
    },
    {
        "name": "tableS10_sif_wavelength",
        "caption": "Table S10. Sensitivity of the SII association to SIF 757 nm and provider-derived SIF 740 nm",
        "landscape": True,
        "build": build_table_s10,
        "inputs": [
            "results/supplementary_outcomes_results.csv",
            "results/supplementary_outcomes_surrogate_summary.csv",
        ],
        "filters": [
            "sample_type = outcome_specific_full",
            "outcome = sif_757nm, sif_740nm",
        ],
        "p_value_source": "results/supplementary_outcomes_surrogate_summary.csv (empirical temporal-surrogate, B = 1000)",
        "bootstrap_source": "N/A",
        "sample_matching": "N/A",
    },
]


def _file_hash(path: Path) -> str:
    """Return SHA-256 hash of a file, or empty string if missing."""
    if not path.exists():
        return ""
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()[:16]


def _input_manifest_hash(specs: list[dict]) -> str:
    """Aggregate hash over all declared input files."""
    h = hashlib.sha256()
    root = project_root()
    for spec in specs:
        for inp in spec["inputs"]:
            p = root / inp
            if p.exists():
                h.update(p.read_bytes())
    return h.hexdigest()[:16]


def _git_commit() -> str:
    """Read short git commit from GIT_COMMIT_INFO.txt if present."""
    info = project_root() / "GIT_COMMIT_INFO.txt"
    if not info.exists():
        return "unknown"
    for line in info.read_text(encoding="utf-8").splitlines():
        if line.startswith("commit:"):
            return line.split(":", 1)[1].strip()
    return "unknown"


def _merge_docx(
    output_path: Path,
    docx_paths: list[Path],
    page_break_before: set[int] | None = None,
) -> None:
    """Merge multiple DOCX files into one, preserving styles.

    ``page_break_before`` contains 0-based indices of documents (relative to the
    merged list) that should start on a new page. The first document cannot have
    a preceding page break.
    """
    if not docx_paths:
        return
    page_break_before = page_break_before or set()
    composer = Composer(Document(str(docx_paths[0])))
    for idx, path in enumerate(docx_paths[1:], start=1):
        if idx in page_break_before:
            # Insert an explicit page break before appending this document.
            from docx.oxml import parse_xml
            p = composer.doc.add_paragraph()
            run = p.add_run()
            run._r.append(
                parse_xml(
                    r'<w:br xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="page"/>'
                )
            )
        composer.append(Document(str(path)))
    composer.save(str(output_path))


def main() -> int:
    root = project_root()
    out_dir = tables_dir()
    git_commit = _git_commit()
    run_timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    input_manifest_hash = _input_manifest_hash(TABLE_SPECS)

    generated: list[dict] = []
    docx_paths_main: list[Path] = []
    docx_paths_supp: list[Path] = []
    supp_page_break_before: set[int] = set()
    supp_idx = 0

    for spec in TABLE_SPECS:
        print(f"[BUILD] {spec['name']}")
        result = spec["build"]()
        if spec["name"] in ("tableS3_fixed_window_results", "tableS6_environmental_driver"):
            display_df, full_df, notes = result
            paths = write_table(
                display_df,
                spec["name"],
                spec["caption"],
                notes=notes,
                landscape=spec["landscape"],
                full_df=full_df,
            )
            row_count = len(display_df)
            source_row_count = len(full_df)
        else:
            display_df, notes = result
            paths = write_table(
                display_df,
                spec["name"],
                spec["caption"],
                notes=notes,
                landscape=spec["landscape"],
            )
            row_count = len(display_df)
            source_row_count = row_count

        docx_path = paths["docx"]
        if spec["name"] == "table1_primary_results":
            docx_paths_main.append(docx_path)
        else:
            if spec["name"] in ("tableS9_lai_diagnostics", "tableS10_sif_wavelength"):
                supp_page_break_before.add(supp_idx)
            docx_paths_supp.append(docx_path)
            supp_idx += 1

        # Verify duplicate analytical keys where applicable.
        duplicate_keys = False
        if spec["name"] == "tableS3_fixed_window_results":
            key_cols = ["Scenario", "Temp. regime", "Method", "Window (d)"]
            if display_df.duplicated(subset=key_cols).any():
                duplicate_keys = True
        elif spec["name"] == "tableS7_shapley":
            key_cols = ["Scenario", "Temperature regime"]
            if display_df.duplicated(subset=key_cols).any():
                duplicate_keys = True

        missing_inputs = [str(root / inp) for inp in spec["inputs"] if not (root / inp).exists()]

        generated.append({
            "name": spec["name"],
            "caption": spec["caption"],
            "csv_path": str(paths["csv"]),
            "md_path": str(paths["md"]),
            "docx_path": str(paths["docx"]),
            "row_count": row_count,
            "source_row_count": source_row_count,
            "duplicate_keys": duplicate_keys,
            "missing_inputs": missing_inputs,
            "inputs": "; ".join(spec["inputs"]),
            "filters": "; ".join(spec["filters"]),
            "p_value_source": spec["p_value_source"],
            "bootstrap_source": spec["bootstrap_source"],
            "sample_matching": spec["sample_matching"],
        })
        print(f"[OK] {spec['name']}: {row_count} display rows ({source_row_count} source rows)")

    # Merge DOCX files.
    main_docx = out_dir / "main_tables.docx"
    supp_docx = out_dir / "supplementary_tables.docx"
    _merge_docx(main_docx, docx_paths_main)
    _merge_docx(supp_docx, docx_paths_supp, page_break_before=supp_page_break_before)
    print(f"[OK] Merged main tables -> {main_docx}")
    print(f"[OK] Merged supplementary tables -> {supp_docx}")

    # Write manifest.
    manifest = pd.DataFrame([
        {
            "table_id": g["name"],
            "caption": g["caption"],
            "csv_path": g["csv_path"],
            "md_path": g["md_path"],
            "docx_path": g["docx_path"],
            "row_count": g["row_count"],
            "source_row_count": g["source_row_count"],
            "primary_inputs": g["inputs"],
            "git_commit": git_commit,
            "input_manifest_hash": input_manifest_hash,
            "status": "blocked" if g["missing_inputs"] or g["duplicate_keys"] else "complete",
        }
        for g in generated
    ])
    manifest_path = out_dir / "tables_manifest.csv"
    manifest.to_csv(manifest_path, index=False)
    print(f"[OK] Manifest -> {manifest_path}")

    # Write validation report.
    validation_lines = [
        "# Tables validation report",
        "",
        f"Generated: {run_timestamp}",
        f"Git commit: {git_commit}",
        f"Input manifest hash: {input_manifest_hash}",
        "",
    ]
    for g in generated:
        validation_lines.extend([
            f"## {g['name']}",
            "",
            f"- Caption: {g['caption']}",
            f"- Display rows: {g['row_count']}",
            f"- Source rows: {g['source_row_count']}",
            f"- CSV: {g['csv_path']}",
            f"- Inputs: {g['inputs']}",
            f"- Filters: {g['filters']}",
            f"- p-value source: {g['p_value_source']}",
            f"- Bootstrap source: {g['bootstrap_source']}",
            f"- Sample matching: {g['sample_matching']}",
            f"- Duplicate analytical keys: {'YES' if g['duplicate_keys'] else 'No'}",
        ])
        if g["missing_inputs"]:
            validation_lines.append(f"- Missing inputs: {', '.join(g['missing_inputs'])}")
        else:
            validation_lines.append("- All declared inputs present: Yes")
        validation_lines.append("")

    # Cross-check with figure source CSVs where applicable.
    validation_lines.extend([
        "## Cross-checks against figure source CSVs",
        "",
    ])
    fig_s3 = root / "reports" / "figures" / "figureS3_surrogate_diagnostics_source.csv"
    fig_s4 = root / "reports" / "figures" / "figureS4_temperature_control_profiles_source.csv"
    fig_s7 = root / "reports" / "figures" / "figureS7_shapley_extended_source.csv"
    for label, path in [
        ("Figure S3 surrogate diagnostics", fig_s3),
        ("Figure S4 temperature control profiles", fig_s4),
        ("Figure S7 Shapley extended", fig_s7),
    ]:
        exists = path.exists()
        validation_lines.append(f"- {label}: {path} ({'present' if exists else 'missing'})")
    validation_lines.append("")

    validation_path = out_dir / "tables_validation.md"
    validation_path.write_text("\n".join(validation_lines), encoding="utf-8")
    print(f"[OK] Validation report -> {validation_path}")

    # Run pytest.
    print("[TEST] Running pytest scripts/tests -q")
    pytest_result = subprocess.run(
        [sys.executable, "-m", "pytest", "scripts/tests", "-q"],
        cwd=root,
        capture_output=True,
        text=True,
    )
    print(pytest_result.stdout)
    if pytest_result.returncode != 0:
        print(pytest_result.stderr, file=sys.stderr)
        print(f"[WARN] pytest exited with code {pytest_result.returncode}")
    else:
        print("[OK] pytest passed")

    # Final summary.
    print("\n=== Table generation summary ===")
    total = len(generated)
    blocked = sum(1 for g in generated if g["missing_inputs"] or g["duplicate_keys"])
    print(f"Tables generated: {total}")
    print(f"Blocked: {blocked}")
    print(f"Main DOCX: {main_docx}")
    print(f"Supplementary DOCX: {supp_docx}")
    print(f"Manifest: {manifest_path}")
    print(f"Validation: {validation_path}")
    return 0 if blocked == 0 and pytest_result.returncode == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
